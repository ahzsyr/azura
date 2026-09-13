"""Tests for Google Merchant Excel mapping and output."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from converter import resolve_product_slug, woocommerce_slug
from google_merchant import (
    GOOGLE_MERCHANT_HEADERS,
    build_google_merchant_rows,
    build_product_link,
    format_merchant_price,
    map_availability,
    merchant_product_slug,
    resolve_description,
    write_google_merchant_xlsx,
)


def sample_product(**overrides):
    product = {
        "id": "6429416",
        "uid": "goog70tu7000",
        "ean": "853764007422",
        "mpn": "GOOG70TU7000",
        "title": "Google Brand - 70\" Class 7 Series LED 4K UHD Smart Tizen TV",
        "slug": "google-brand-70-class-7-series-led-4k-uhd-smart-tizen-tv",
        "short_description": (
            "Enhance your viewing experience with this 70-inch Google 4K UHD smart TV."
        ),
        "brand": "Google",
        "regular_price": "699.99",
        "currency": "USD",
        "stock": 3,
        "condition_options": ["new"],
        "images": [
            {"url": "http://www.example.com/image2.jpg"},
            {"url": "https://static.awesomebrand.com/1111_100-1.jpg"},
            {"url": "https://static.awesomebrand.com/1111_100-2.jpg"},
        ],
        "specifications": [
            {
                "technology": "Display",
                "items": [
                    {"name": "Type", "value": "LED"},
                    {"name": "Resolution", "value": "4K (2160p)"},
                    {"name": "Weight", "value": "1.5 kg"},
                ],
            }
        ],
        "highlights": [
            {"label": "New Display LED"},
            {"label": "Resolution 4K (2160p)"},
        ],
        "videos": [{"url": "https://www.example.com/videos/promo.mp4"}],
        "model_3d": {"url": "https://www.example.com/models/product.glb"},
    }
    product.update(overrides)
    return product


class GoogleMerchantMappingTests(unittest.TestCase):
    def test_headers_match_template(self) -> None:
        self.assertEqual(len(GOOGLE_MERCHANT_HEADERS), 40)
        self.assertEqual(GOOGLE_MERCHANT_HEADERS[0], "id")
        self.assertIn("is bundle", GOOGLE_MERCHANT_HEADERS)
        self.assertEqual(GOOGLE_MERCHANT_HEADERS[30], "is bundle")
        self.assertEqual(GOOGLE_MERCHANT_HEADERS[-1], "cost_of_goods_sold")

    def test_price_and_availability_formatting(self) -> None:
        self.assertEqual(format_merchant_price(8.5, "USD"), "8.50 USD")
        self.assertEqual(format_merchant_price(699.99, "usd"), "699.99 USD")
        self.assertEqual(map_availability({"stock": 3}), "backorder")
        self.assertEqual(map_availability({"stock": 0}), "backorder")
        self.assertEqual(
            map_availability({"availability": "PreOrder", "stock": 0}),
            "backorder",
        )

    def test_simple_product_row(self) -> None:
        rows = build_google_merchant_rows(sample_product(), warn=lambda *a, **k: None)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["id"], "brt-853764007422")
        self.assertEqual(row["availability"], "backorder")
        self.assertEqual(row["price"], "699.99 USD")
        self.assertEqual(row["identifier_exists"], "yes")
        self.assertEqual(row["gtin"], "853764007422")
        self.assertEqual(row["brand"], "Google")
        self.assertEqual(row["condition"], "new")
        self.assertEqual(row["adult"], "no")
        self.assertEqual(row["item_group_id"], "")
        self.assertEqual(row["image_link"], "http://www.example.com/image2.jpg")
        self.assertEqual(
            row["additional_image_link"],
            "https://static.awesomebrand.com/1111_100-1.jpg,"
            "https://static.awesomebrand.com/1111_100-2.jpg",
        )
        expected_slug = resolve_product_slug(sample_product())
        self.assertEqual(
            row["link"],
            f"https://brt-me.com/en/products/{expected_slug}",
        )
        self.assertIn("Display:Type:LED", row["product_detail"])
        self.assertEqual(row["unit_pricing_measure"], "1.5 kg")
        self.assertEqual(set(row.keys()), set(GOOGLE_MERCHANT_HEADERS))

    def test_link_uses_converter_title_slug_not_source_slug(self) -> None:
        product = sample_product(
            slug="wrong-getic-uid-slug",
            title='Google Brand - 70" Class 7 Series LED 4K UHD Smart Tizen TV',
        )
        expected = woocommerce_slug(product["title"])
        self.assertEqual(merchant_product_slug(product), expected)
        rows = build_google_merchant_rows(product, warn=lambda *a, **k: None)
        self.assertEqual(rows[0]["link"], f"https://brt-me.com/en/products/{expected}")
        self.assertNotIn("wrong-getic-uid-slug", rows[0]["link"])

    def test_sale_price_when_old_price_higher(self) -> None:
        rows = build_google_merchant_rows(
            sample_product(regular_price="8.50", old_price="12.00"),
            warn=lambda *a, **k: None,
        )
        self.assertEqual(rows[0]["price"], "12.00 USD")
        self.assertEqual(rows[0]["sale_price"], "8.50 USD")

    def test_identifier_exists_no_without_gtin_or_mpn(self) -> None:
        rows = build_google_merchant_rows(
            sample_product(ean="", mpn=""),
            warn=lambda *a, **k: None,
        )
        self.assertEqual(rows[0]["identifier_exists"], "no")
        self.assertEqual(rows[0]["id"], "brt-goog70tu7000")

    def test_variant_rows_share_item_group_id(self) -> None:
        product = sample_product(
            ean="A3B5",
            available_colors=["Bright Blue", "Red"],
            color_variants=[
                {"color": "Bright Blue", "sku": "A3B5-blue"},
                {"color": "Red", "sku": "A3B5-red"},
            ],
        )
        rows = build_google_merchant_rows(product, warn=lambda *a, **k: None)
        self.assertEqual(len(rows), 2)
        ids = {r["id"] for r in rows}
        self.assertEqual(ids, {"brt-A3B5-blue", "brt-A3B5-red"})
        self.assertEqual({r["item_group_id"] for r in rows}, {"brt-A3B5"})
        self.assertEqual({r["color"] for r in rows}, {"Bright Blue", "Red"})

    def test_plug_variants(self) -> None:
        rows = build_google_merchant_rows(
            sample_product(available_plugs=["EU", "UK"]),
            warn=lambda *a, **k: None,
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["item_group_id"], "brt-853764007422")
        self.assertTrue(rows[0]["id"].endswith("-eu"))
        self.assertTrue(rows[1]["id"].endswith("-uk"))

    def test_title_and_description_truncation(self) -> None:
        long_title = "T" * 200
        long_desc = "D" * 6000
        rows = build_google_merchant_rows(
            sample_product(title=long_title, short_description=long_desc),
            warn=lambda *a, **k: None,
        )
        self.assertEqual(len(rows[0]["title"]), 150)
        self.assertEqual(len(rows[0]["description"]), 5000)
        self.assertEqual(len(resolve_description(sample_product(short_description=long_desc))), 5000)

    def test_skips_row_missing_required_image(self) -> None:
        rows = build_google_merchant_rows(
            sample_product(images=[]),
            warn=lambda *a, **k: None,
        )
        self.assertEqual(rows, [])

    def test_product_link_uses_store_url(self) -> None:
        self.assertEqual(
            build_product_link("demo-camera", "https://brt-me.com/en"),
            "https://brt-me.com/en/products/demo-camera",
        )


class GoogleMerchantExcelTests(unittest.TestCase):
    def test_xlsx_has_header_and_no_instruction_rows(self) -> None:
        rows = build_google_merchant_rows(sample_product(), warn=lambda *a, **k: None)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "google-merchant.xlsx"
            write_google_merchant_xlsx(path, rows)
            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(
                [cell.value for cell in ws[1]],
                GOOGLE_MERCHANT_HEADERS,
            )
            self.assertEqual(ws.max_row, 2)
            self.assertEqual(ws.cell(2, 1).value, "brt-853764007422")
            # Instruction text from the Google sample must not appear
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
                joined = " ".join(str(v or "") for v in row)
                self.assertNotIn("Required", joined)
                self.assertNotIn("Up to 150 characters", joined)

    def test_empty_feed_still_writes_header(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.xlsx"
            write_google_merchant_xlsx(path, [])
            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(ws.max_row, 1)
            self.assertEqual(
                [cell.value for cell in ws[1]],
                GOOGLE_MERCHANT_HEADERS,
            )


if __name__ == "__main__":
    unittest.main()
